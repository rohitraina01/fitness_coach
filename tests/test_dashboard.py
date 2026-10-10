import os
import unittest
import openpyxl

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DASHBOARD_PATH = os.path.join(BASE_DIR, "DASHBOARD", "fitness_dashboard.xlsx")

class TestDashboard(unittest.TestCase):
    def test_dashboard_file_exists(self):
        self.assertTrue(os.path.exists(DASHBOARD_PATH))
        self.assertGreater(os.path.getsize(DASHBOARD_PATH), 5000)

    def test_dashboard_worksheets(self):
        wb = openpyxl.load_workbook(DASHBOARD_PATH)
        sheet_names = wb.sheetnames
        expected_sheets = [
            "Executive_Dashboard",
            "Workouts",
            "Workout_Sets",
            "Exercise_Catalog",
            "Body_Composition",
            "Nutrition_Activity",
            "Volume_Attribution",
            "Weekly_Adherence"
        ]
        for s in expected_sheets:
            self.assertIn(s, sheet_names, f"Missing sheet: {s}")

    def test_structured_tables_present(self):
        wb = openpyxl.load_workbook(DASHBOARD_PATH)
        ws_workouts = wb["Workouts"]
        self.assertIn("Table_Workouts", ws_workouts.tables)
        ws_sets = wb["Workout_Sets"]
        self.assertIn("Table_WorkoutSets", ws_sets.tables)

if __name__ == "__main__":
    unittest.main()
